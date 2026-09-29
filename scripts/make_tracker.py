import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, Reference

def create_tracker():
    wb = openpyxl.Workbook()
    
    # ---------------------------
    # SHEET 1: Applications
    # ---------------------------
    ws1 = wb.active
    ws1.title = "Applications"
    
    columns = [
        "Company", "Position", "Website", "Link to posting", "Date found", 
        "Source", "Location / Work mode", "Address", "Google Maps", 
        "Commute / travel time", "Contact person", "Contact e-mail", 
        "Referral / Advocate", "What interests me in the offer", "Role fit score", 
        "Priority", "Tailored application?", "Documents used", "Portfolio / work sample link", 
        "Deadline", "Date sent", "Follow-up date", "Last contact date", 
        "Status", "Answer", "Interview date", "Interview type", 
        "Salary / Stipend", "Next action", "Next action date", "Reason lost", 
        "Notes", "Duplicate check"
    ]
    
    ws1.append(columns)
    
    # Header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    for cell in ws1[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        
    # Freeze top row and first column (Company)
    ws1.freeze_panes = "B2"
    
    # Add auto-filter
    ws1.auto_filter.ref = f"A1:AG200"
    
    # Column widths & wrapping
    ws1.column_dimensions['A'].width = 20  # Company
    ws1.column_dimensions['B'].width = 20  # Position
    ws1.column_dimensions['N'].width = 30  # What interests me
    ws1.column_dimensions['AF'].width = 30 # Notes
    
    for row_idx in range(2, 202):
        ws1.cell(row=row_idx, column=14).alignment = Alignment(wrap_text=True) # What interests me
        ws1.cell(row=row_idx, column=32).alignment = Alignment(wrap_text=True) # Notes

        # Formulas
        # I: Google Maps
        ws1.cell(row=row_idx, column=9).value = f'=IF(H{row_idx}="","",HYPERLINK("https://www.google.com/maps/search/?api=1&query="&SUBSTITUTE(A{row_idx}&" "&H{row_idx}," ","+"),"Open map"))'
        # V: Follow-up date
        ws1.cell(row=row_idx, column=22).value = f'=IF(ISBLANK(U{row_idx}),"",U{row_idx}+7)'
        # AG: Duplicate check
        ws1.cell(row=row_idx, column=33).value = f'=IF(COUNTIF(A:A,A{row_idx})>1, IF(COUNTIFS(A:A,A{row_idx},B:B,B{row_idx})>1,"DUPLICATE",""),"")'
        
        # Date formatting for columns 5 (E), 20 (T), 21 (U), 22 (V), 23 (W), 26 (Z), 30 (AD)
        for c in [5, 20, 21, 22, 23, 26, 30]:
            ws1.cell(row=row_idx, column=c).number_format = 'DD/MM/YYYY'

    # Data Validations
    dv_source = DataValidation(type="list", formula1='"LinkedIn,Company site,Referral,Job fair,Other"', allow_blank=True)
    dv_work_mode = DataValidation(type="list", formula1='"On-site,Hybrid,Remote"', allow_blank=True)
    dv_fit = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    dv_priority = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    dv_tailored = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv_status = DataValidation(type="list", formula1='"To apply,Applied,Follow-up sent,Interview scheduled,Offer,Rejected,Ghosted"', allow_blank=True)
    dv_answer = DataValidation(type="list", formula1='"No answer,Positive,Negative"', allow_blank=True)
    dv_interview_type = DataValidation(type="list", formula1='"Phone screen,Technical,Case study,Hiring manager,Final round"', allow_blank=True)
    dv_reason = DataValidation(type="list", formula1='"No answer,Skills gap,Salary mismatch,Position closed,Location,I declined,Other"', allow_blank=True)
    
    ws1.add_data_validation(dv_source)
    ws1.add_data_validation(dv_work_mode)
    ws1.add_data_validation(dv_fit)
    ws1.add_data_validation(dv_priority)
    ws1.add_data_validation(dv_tailored)
    ws1.add_data_validation(dv_status)
    ws1.add_data_validation(dv_answer)
    ws1.add_data_validation(dv_interview_type)
    ws1.add_data_validation(dv_reason)

    dv_source.add("F2:F200")
    dv_work_mode.add("G2:G200")
    dv_fit.add("O2:O200")
    dv_priority.add("P2:P200")
    dv_tailored.add("Q2:Q200")
    dv_status.add("X2:X200")
    dv_answer.add("Y2:Y200")
    dv_interview_type.add("AA2:AA200")
    dv_reason.add("AE2:AE200")

    # Conditional formatting
    red_fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
    orange_fill = PatternFill(start_color='FFE5CC', end_color='FFE5CC', fill_type='solid')
    green_fill = PatternFill(start_color='CCFFCC', end_color='CCFFCC', fill_type='solid')
    grey_font = Font(color='808080')
    red_font = Font(color='FF0000', bold=True)
    
    # Red fill: Next action date or Follow-up date has passed, Status "Applied", Answer "No answer"
    rule_red = FormulaRule(formula=['AND($X2="Applied", $Y2="No answer", OR(AND(ISNUMBER($V2), $V2<TODAY()), AND(ISNUMBER($AD2), $AD2<TODAY())))'], stopIfTrue=True, fill=red_fill)
    ws1.conditional_formatting.add("A2:AG200", rule_red)
    
    # Orange fill: Deadline within next 3 days and Status "To apply"
    rule_orange = FormulaRule(formula=['AND($X2="To apply", ISNUMBER($T2), $T2-TODAY()<=3, $T2-TODAY()>=0)'], stopIfTrue=True, fill=orange_fill)
    ws1.conditional_formatting.add("A2:AG200", rule_orange)
    
    # Green fill: Status "Interview scheduled" or "Offer"
    rule_green = FormulaRule(formula=['OR($X2="Interview scheduled", $X2="Offer")'], stopIfTrue=True, fill=green_fill)
    ws1.conditional_formatting.add("A2:AG200", rule_green)
    
    # Grey text: Status "Rejected" or "Ghosted"
    rule_grey = FormulaRule(formula=['OR($X2="Rejected", $X2="Ghosted")'], stopIfTrue=True, font=grey_font)
    ws1.conditional_formatting.add("A2:AG200", rule_grey)
    
    # Red text for DUPLICATE in Duplicate check
    rule_duplicate = CellIsRule(operator='equal', formula=['"DUPLICATE"'], stopIfTrue=True, font=red_font)
    ws1.conditional_formatting.add("AG2:AG200", rule_duplicate)

    # Example rows
    ws1.append(["Example Corp", "Software Engineer", "example.com", "example.com/job", "01/01/2026", "LinkedIn", "Remote", "123 Tech St", "", "N/A", "John Doe", "john@example.com", "", "Tech stack", 5, "High", "Yes", "CV_v1", "portfolio.com", "15/01/2026", "05/01/2026", "", "", "Applied", "No answer", "", "", "$100k", "Follow up", "12/01/2026", "", "DELETE THIS EXAMPLE ROW", ""])
    ws1.append(["Acme Inc", "Data Analyst", "acme.com", "acme.com/job", "02/01/2026", "Company site", "Hybrid", "456 Data Ave", "", "30 mins", "", "", "Jane Doe", "Growth", 4, "Medium", "No", "CV_v2", "", "20/01/2026", "", "", "", "To apply", "", "", "", "$80k", "Draft CV", "", "", "DELETE THIS EXAMPLE ROW", ""])
    
    # Format example row dates
    for r in [2, 3]:
        for c in [5, 20, 21, 22, 23, 26, 30]:
            if isinstance(ws1.cell(row=r, column=c).value, str) and "/" in ws1.cell(row=r, column=c).value:
                parts = ws1.cell(row=r, column=c).value.split("/")
                if len(parts) == 3:
                    import datetime
                    ws1.cell(row=r, column=c).value = datetime.datetime(int(parts[2]), int(parts[1]), int(parts[0]))
            ws1.cell(row=r, column=c).number_format = 'DD/MM/YYYY'


    # ---------------------------
    # SHEET 2: Summary
    # ---------------------------
    ws2 = wb.create_sheet("Summary")
    
    ws2.append(["Metric", "Value"])
    ws2.cell(row=1, column=1).font = Font(bold=True)
    ws2.cell(row=1, column=2).font = Font(bold=True)
    
    ws2.append(["Total applications", "=COUNTA(Applications!A2:A200)"])
    ws2.append(["Applications sent", '=COUNTIFS(Applications!X2:X200,"<>To apply",Applications!X2:X200,"<>")'])
    ws2.append(["Response rate", '=IF(B3=0,"0%",TEXT(COUNTIFS(Applications!Y2:Y200,"<>No answer",Applications!Y2:Y200,"<>",Applications!Y2:Y200,"<>""") / B3, "0.0%"))'])
    ws2.append(["Interview rate", '=IF(B3=0,"0%",TEXT(COUNTIF(Applications!X2:X200,"Interview scheduled") / B3, "0.0%"))'])
    ws2.append(["Number of offers", '=COUNTIF(Applications!X2:X200,"Offer")'])
    
    ws2.append([])
    ws2.append(["Applications by Status", "Count"])
    ws2.cell(row=8, column=1).font = Font(bold=True)
    
    statuses = ["To apply", "Applied", "Follow-up sent", "Interview scheduled", "Offer", "Rejected", "Ghosted"]
    for i, st in enumerate(statuses, start=9):
        ws2.append([st, f'=COUNTIF(Applications!X2:X200,"{st}")'])
        
    ws2.append([])
    ws2.append(["Applications by Source", "Count"])
    ws2.cell(row=17, column=1).font = Font(bold=True)
    
    sources = ["LinkedIn", "Company site", "Referral", "Job fair", "Other"]
    for i, src in enumerate(sources, start=18):
        ws2.append([src, f'=COUNTIF(Applications!F2:F200,"{src}")'])
        
    # Bar Chart: Status
    chart1 = BarChart()
    chart1.title = "Applications by Status"
    data1 = Reference(ws2, min_col=2, min_row=8, max_row=15)
    cats1 = Reference(ws2, min_col=1, min_row=9, max_row=15)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    ws2.add_chart(chart1, "D2")
    
    # Bar Chart: Source
    chart2 = BarChart()
    chart2.title = "Applications by Source"
    data2 = Reference(ws2, min_col=2, min_row=17, max_row=22)
    cats2 = Reference(ws2, min_col=1, min_row=18, max_row=22)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    ws2.add_chart(chart2, "D18")
    
    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 15
    
    wb.save("job_search_tracker.xlsx")
    print("job_search_tracker.xlsx created successfully.")

if __name__ == "__main__":
    create_tracker()
